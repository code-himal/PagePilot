"""Turn an (optionally user-edited) extraction result into DOCX / TXT / XLSX / CSV files.

Input is untrusted JSON posted by the browser, so `normalize_document` rebuilds a clean structure
and enforces size limits before anything is written.
"""
import csv
import io
import re
import zipfile
from typing import Dict, List, Optional, Tuple

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import Config
from .errors import ExtractError
from .numbers import parse_number
from .textutil import clean_text, safe_filename

MIME = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "txt": "text/plain; charset=utf-8",
    "csv": "text/csv; charset=utf-8",
    "zip": "application/zip",
}
_MAX_TEXT = 32000                    # Excel's cell limit is 32767 characters


# ----------------------------------------------------------------------------- validation
def _bad(msg: str):
    raise ExtractError("bad_request", msg, 400)


def normalize_document(raw) -> dict:
    """Validate and clean one document from the browser."""
    if not isinstance(raw, dict):
        _bad("Invalid document.")
    pages_raw = raw.get("pages")
    if not isinstance(pages_raw, list) or len(pages_raw) > 500:
        _bad("Invalid document pages.")

    cells_total = 0
    pages = []
    for p in pages_raw:
        if not isinstance(p, dict) or not isinstance(p.get("blocks", []), list):
            _bad("Invalid page.")
        number = p.get("number")
        number = number if isinstance(number, int) and 0 < number < 100000 else len(pages) + 1
        blocks, tcount = [], 0
        for b in p.get("blocks", []):
            if not isinstance(b, dict):
                _bad("Invalid block.")
            t = b.get("type")
            if t in ("heading", "paragraph", "list_item"):
                text = clean_text(b.get("text", ""))[:_MAX_TEXT]
                block = {"type": t, "text": text}
                if t == "heading":
                    lvl = b.get("level")
                    block["level"] = lvl if lvl in (1, 2, 3) else 2
                blocks.append(block)
            elif t == "table":
                rows_raw = b.get("rows")
                if not isinstance(rows_raw, list) or len(rows_raw) > 5000:
                    _bad("Invalid table.")
                rows = []
                for r in rows_raw:
                    if not isinstance(r, list) or len(r) > 200:
                        _bad("Invalid table row.")
                    rows.append([clean_text(c)[:_MAX_TEXT] for c in r])
                width = max((len(r) for r in rows), default=0)
                rows = [r + [""] * (width - len(r)) for r in rows]
                cells_total += width * len(rows)
                if cells_total > Config.MAX_EXPORT_CELLS:
                    _bad("This document is too large to export.")
                tcount += 1
                blocks.append({"type": "table", "rows": rows, "header": bool(b.get("header", True)),
                               "id": f"p{number}t{tcount}"})
        pages.append({"number": number, "blocks": blocks})

    fields = []
    for f in (raw.get("fields") or [])[:500]:
        if isinstance(f, dict):
            fields.append({"key": clean_text(f.get("key", ""))[:200],
                           "value": clean_text(f.get("value", ""))[:2000],
                           "page": f.get("page") if isinstance(f.get("page"), int) else ""})
    return {"filename": safe_filename(str(raw.get("filename") or "document")), "pages": pages, "fields": fields}


def list_tables(doc: dict) -> List[dict]:
    out = []
    for pg in doc["pages"]:
        for b in pg["blocks"]:
            if b["type"] == "table" and b["rows"]:
                out.append({"id": b["id"], "page": pg["number"], "rows": b["rows"], "header": b["header"]})
    return out


# ----------------------------------------------------------------------------- TXT
def to_txt(doc: dict) -> bytes:
    parts = []
    multi = len([p for p in doc["pages"] if p["blocks"]]) > 1
    for pg in doc["pages"]:
        if not pg["blocks"]:
            continue
        if multi:
            parts.append(f"----- Page {pg['number']} -----")
        for b in pg["blocks"]:
            if b["type"] == "heading":
                parts.append(b["text"])
            elif b["type"] == "paragraph":
                parts.append(b["text"])
            elif b["type"] == "list_item":
                parts.append("- " + b["text"])
            elif b["type"] == "table":
                parts.append("\n".join("\t".join(c.replace("\t", " ").replace("\n", " ") for c in r)
                                       for r in b["rows"]))
    text = "\n\n".join(p for p in parts if p != "")
    return (text + "\n").encode("utf-8")


# ----------------------------------------------------------------------------- DOCX
def _shade(cell, hex_fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _add_text_with_breaks(paragraph, text: str, size: Optional[float] = None, bold: bool = False):
    lines = text.split("\n")
    for i, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.bold = bold or None
        if size:
            run.font.size = Pt(size)
        if i < len(lines) - 1:
            run.add_break()


def to_docx(doc: dict) -> bytes:
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)          # A4
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.2))

    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(6)
    for i, size in ((1, 20), (2, 15), (3, 12.5)):
        st = d.styles[f"Heading {i}"]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)

    d.core_properties.title = doc["filename"]
    d.core_properties.author = "DocPixly"

    wrote_any = False
    first_page = True
    for pg in doc["pages"]:
        if not pg["blocks"]:
            continue
        if not first_page:
            d.add_page_break()
        first_page = False
        for b in pg["blocks"]:
            wrote_any = True
            if b["type"] == "heading":
                d.add_heading(b["text"], level=b["level"])
            elif b["type"] == "paragraph":
                _add_text_with_breaks(d.add_paragraph(), b["text"])
            elif b["type"] == "list_item":
                _add_text_with_breaks(d.add_paragraph(style="List Bullet"), b["text"])
            elif b["type"] == "table" and b["rows"]:
                _docx_table(d, b)
    if not wrote_any:
        d.add_paragraph("No text was found in this file.")

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _docx_table(d, b):
    rows = b["rows"]
    ncols = len(rows[0])
    table = d.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    size = 10 if ncols <= 5 else 9 if ncols <= 8 else 7.5
    for i, row in enumerate(rows):
        is_header = b["header"] and i == 0
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            _add_text_with_breaks(para, value, size=size, bold=is_header)
            if is_header:
                _shade(cell, "E4E8F5")
        if is_header:
            _repeat_header(table.rows[0])
    d.add_paragraph().paragraph_format.space_after = Pt(4)


# ----------------------------------------------------------------------------- XLSX
_SHEET_BAD = re.compile(r"[\[\]:*?/\\]")


def _sheet_name(base: str, used: set) -> str:
    name = _SHEET_BAD.sub(" ", base).strip()[:31] or "Sheet"
    candidate, n = name, 2
    while candidate.lower() in used:
        suffix = f" ({n})"
        candidate = name[:31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def to_xlsx(doc: dict, numbers: bool = True) -> bytes:
    tables = list_tables(doc)
    if not tables:
        raise ExtractError("no_tables", "No tables were found in this file, so there is nothing to put in Excel.", 422)

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = doc["filename"]
    wb.properties.creator = "DocPixly"
    used: set = set()
    base_font = Font(name="Arial", size=10)
    head_font = Font(name="Arial", size=10, bold=True)
    head_fill = PatternFill("solid", start_color="E4E8F5", end_color="E4E8F5")
    thin = Side(style="thin", color="8A93B2")

    for t in tables:
        ws = wb.create_sheet(_sheet_name(f"Page {t['page']} Table {t['id'].split('t')[-1]}", used))
        widths: Dict[int, int] = {}
        for i, row in enumerate(t["rows"], start=1):
            for j, value in enumerate(row, start=1):
                cell = ws.cell(row=i, column=j)
                parsed = parse_number(value) if (numbers and value and not (t["header"] and i == 1)) else None
                if parsed:
                    cell.value, fmt = parsed
                    if fmt:
                        cell.number_format = fmt
                    shown = len(value)
                else:
                    cell.value = value
                    if isinstance(value, str) and value.startswith("="):
                        cell.data_type = "s"                    # text that merely starts with '=' is not a formula
                    shown = max((len(x) for x in value.split("\n")), default=0)
                widths[j] = max(widths.get(j, 0), shown)
                cell.font = head_font if (t["header"] and i == 1) else base_font
                if t["header"] and i == 1:
                    cell.fill = head_fill
                    cell.border = Border(bottom=thin)
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                elif "\n" in value:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        for j, w in widths.items():
            ws.column_dimensions[get_column_letter(j)].width = min(max(w * 1.1 + 2, 8), 60)
        if t["header"] and len(t["rows"]) > 1:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(len(t['rows'][0]))}{len(t['rows'])}"

    if doc["fields"]:
        ws = wb.create_sheet(_sheet_name("Fields", used))
        ws.append(["Field", "Value", "Page"])
        for f in doc["fields"]:
            ws.append([f["key"], f["value"], f["page"]])
        for row in ws.iter_rows():
            for cell in row:
                cell.font = head_font if cell.row == 1 else base_font
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"
        ws["A1"].fill = ws["B1"].fill = ws["C1"].fill = head_fill
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 50
        ws.column_dimensions["C"].width = 8
        ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------- CSV
def _csv_safe(value: str) -> str:
    """Stop spreadsheet apps from running cell text as a formula (CSV injection)."""
    if not value:
        return value
    if value[0] in "=+@\t\r":
        return "'" + value
    if value[0] == "-" and parse_number(value) is None:
        return "'" + value
    return value


def to_csv(rows: List[List[str]]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    for r in rows:
        w.writerow([_csv_safe(c) for c in r])
    return buf.getvalue().encode("utf-8-sig")             # BOM so Excel reads non-English text correctly


# ----------------------------------------------------------------------------- entry point
def _unique(name: str, taken: set) -> str:
    base, ext = name.rsplit(".", 1)
    candidate, n = name, 2
    while candidate.lower() in taken:
        candidate = f"{base} ({n}).{ext}"
        n += 1
    taken.add(candidate.lower())
    return candidate


def _zip(files: List[Tuple[str, bytes]]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files:
            z.writestr(name, data)
    return buf.getvalue()


def export(docs_raw: list, fmt: str, numbers: bool = True, table_id: Optional[str] = None,
           doc_index: int = 0) -> Tuple[bytes, str, str]:
    """Returns (file bytes, download filename, mime type)."""
    if fmt not in ("docx", "txt", "xlsx", "csv"):
        _bad("Unknown export format.")
    if not isinstance(docs_raw, list) or not docs_raw or len(docs_raw) > 50:
        _bad("Nothing to export.")
    docs = [normalize_document(d) for d in docs_raw]

    if fmt == "csv":
        return _export_csv(docs, table_id, doc_index)

    taken: set = set()
    files: List[Tuple[str, bytes]] = []
    skipped_no_tables = 0
    for d in docs:
        if fmt == "docx":
            data = to_docx(d)
        elif fmt == "txt":
            data = to_txt(d)
        else:
            if not list_tables(d):
                skipped_no_tables += 1
                continue
            data = to_xlsx(d, numbers)
        files.append((_unique(f"{d['filename']}.{fmt}", taken), data))

    if not files:
        raise ExtractError("no_tables", "No tables were found, so there is nothing to put in Excel.", 422)
    if len(files) == 1:
        return files[0][1], files[0][0], MIME[fmt]
    return _zip(files), f"docpixly-{fmt}.zip", MIME["zip"]


def _export_csv(docs: List[dict], table_id: Optional[str], doc_index: int):
    if table_id:
        d = docs[doc_index if 0 <= doc_index < len(docs) else 0]
        for t in list_tables(d):
            if t["id"] == table_id:
                return to_csv(t["rows"]), f"{d['filename']}-{table_id}.csv", MIME["csv"]
        raise ExtractError("no_tables", "That table was not found.", 404)

    per_doc = [(d, list_tables(d)) for d in docs]
    if not any(ts for _, ts in per_doc):
        raise ExtractError("no_tables", "No tables were found in this file, so there is nothing to save as CSV.", 422)
    if len(docs) == 1 and len(per_doc[0][1]) == 1:
        d, ts = per_doc[0]
        return to_csv(ts[0]["rows"]), f"{d['filename']}.csv", MIME["csv"]

    files, taken = [], set()
    for d, ts in per_doc:
        prefix = f"{d['filename']}/" if len(docs) > 1 else ""
        for t in ts:
            name = f"{prefix}{d['filename']}-page{t['page']}-table{t['id'].split('t')[-1]}.csv"
            files.append((_unique(name, taken), to_csv(t["rows"])))
    return _zip(files), f"{docs[0]['filename'] if len(docs) == 1 else 'docpixly'}-csv.zip", MIME["zip"]
